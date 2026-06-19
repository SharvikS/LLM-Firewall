"""Tests for file/image attachment scanning — analyzer.extract + scan_file.

OCR (image) is covered separately and only when an OCR backend is installed, so
the suite stays green on a minimal environment. Document extractors (PDF/DOCX/
XLSX) are skipped individually when their optional library is absent.
"""

import importlib.util
import io
import json

import pytest

from analyzer import extract
from analyzer.server import AnalyzerServicer


@pytest.fixture(scope="module")
def servicer():
    return AnalyzerServicer()


def _have(mod):
    return importlib.util.find_spec(mod) is not None


# ── extract.py ────────────────────────────────────────────────────────────────

def test_plain_text_extracts_without_any_library():
    r = extract.extract_text("notes.txt", "text/plain", b"hello world")
    assert r.supported and r.kind == "file" and r.text == "hello world"


def test_code_file_is_text():
    r = extract.extract_text("main.py", "", b"api_key = 'x'\n")
    assert r.supported and r.extractor == "text"


def test_unsupported_binary_degrades_gracefully():
    r = extract.extract_text("blob.bin", "application/octet-stream", b"\x00\x01\x02\xff\xfe")
    assert not r.supported and r.error


def test_image_without_backend_is_unsupported_not_crash():
    # Whatever the environment, an undecodable image must not raise.
    r = extract.extract_text("x.png", "image/png", b"\x89PNG\r\n\x1a\n not-real")
    assert r.kind == "image"
    # supported depends on OCR availability; the contract is "no exception".


def test_spoofed_image_content_type_is_not_treated_as_image():
    # A non-image that lies about its content_type must NOT get the privileged
    # (fail-open) "image" classification — the bytes decide, not the label.
    r = extract.extract_text("evil.exe", "image/png", b"\x4d\x5a\x90\x00\x03")  # "MZ" = PE
    assert r.kind == "file" and not r.supported


def test_real_image_mislabelled_as_binary_is_still_an_image():
    # Conversely, real image bytes mislabelled as octet-stream are recognised.
    r = extract.extract_text("photo.bin", "application/octet-stream", b"\x89PNG\r\n\x1a\n x")
    assert r.kind == "image"


def test_sniff_image_covers_common_formats():
    assert extract._sniff_image(b"\xff\xd8\xff\xe0 jfif")      # JPEG
    assert extract._sniff_image(b"GIF89a....")                 # GIF
    assert extract._sniff_image(b"RIFF\x00\x00\x00\x00WEBPVP")  # WEBP
    assert not extract._sniff_image(b"%PDF-1.7")               # not an image
    assert not extract._sniff_image(b"")


def test_spoofed_image_fails_closed_in_strict(servicer):
    # The spoofing payoff would be a strict-mode bypass; confirm it's denied.
    v = servicer.scan_file("evil.exe", "image/png", b"\x4d\x5a\x90\x00", strict=True)
    assert v["decision"] == "block" and v["kind"] == "file"


@pytest.mark.skipif(not _have("docx"), reason="python-docx not installed")
def test_docx_extraction():
    import docx
    d = docx.Document()
    d.add_paragraph("contact bob@example.com")
    buf = io.BytesIO()
    d.save(buf)
    r = extract.extract_text("doc.docx", "", buf.getvalue())
    assert r.supported and "bob@example.com" in r.text


@pytest.mark.skipif(not _have("openpyxl"), reason="openpyxl not installed")
def test_xlsx_extraction():
    import openpyxl
    wb = openpyxl.Workbook()
    wb.active["A1"] = "ssn 123-45-6789"
    buf = io.BytesIO()
    wb.save(buf)
    r = extract.extract_text("sheet.xlsx", "", buf.getvalue())
    assert r.supported and "123-45-6789" in r.text


# ── scan_file (verdict shaping) ────────────────────────────────────────────────

def test_clean_text_file_allows(servicer):
    v = servicer.scan_file("readme.txt", "text/plain", b"The capital of France is Paris.")
    assert v["decision"] == "allow"
    assert v["kind"] == "file" and v["filename"] == "readme.txt"


def test_secret_in_file_blocks(servicer):
    # A file finding can't be redacted in place → it becomes a hard block.
    # (A bare secret like this reads as a secret; some phrasings also trip the
    # injection classifier — either way the contract is: not allowed through.)
    v = servicer.scan_file("creds.txt", "text/plain",
                           b"AKIAIOSFODNN7EXAMPLE")
    assert v["decision"] == "block"
    assert v["categories"]  # something was flagged
    assert v["kind"] == "file"


def test_pii_in_file_is_blocked_not_redacted(servicer):
    # PII alone would be "redact" for text; for a file it escalates to block.
    v = servicer.scan_file("contacts.csv", "text/csv", b"name,email\nBob,bob@example.com\n")
    assert v["decision"] == "block"
    assert "pii" in v["categories"]


def test_unscannable_file_blocks_only_in_strict(servicer):
    bad = b"\x00\x01\x02\xff"
    assert servicer.scan_file("x.bin", "application/octet-stream", bad, strict=True)["decision"] == "block"
    assert servicer.scan_file("x.bin", "application/octet-stream", bad, strict=False)["decision"] == "allow"


def test_unscannable_image_never_blocks_even_in_strict(servicer):
    # An image that can't be OCR'd (or with no readable text) always fails OPEN —
    # it's only ever blocked on a positive finding from a completed scan.
    bad_img = b"\x89PNG\r\n\x1a\n not-a-real-image"
    assert servicer.scan_file("shot.png", "image/png", bad_img, strict=True)["decision"] == "allow"
    assert servicer.scan_file("shot.png", "image/png", bad_img, strict=False)["decision"] == "allow"


def test_file_verdict_is_json_serialisable(servicer):
    v = servicer.scan_file("creds.txt", "text/plain", b"email a@b.com")
    json.dumps(v)  # must not raise
